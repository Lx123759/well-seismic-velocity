"""按模板填写《软件工程实践》开题任务书，输出到项目根目录，不修改原模板。"""
from copy import deepcopy
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
FIG = Path(__file__).resolve().parent / "fig"
TEMPLATE = ROOT / "软件工程实践_开题任务书_软件开发类.docx"
OUTPUT = ROOT / "开题任务书_测井约束下基于深度学习的井震联合速度场建模方法与系统实现.docx"

TITLE = "测井约束下基于深度学习的井震联合速度场建模方法与系统实现"
BODY_SIZE = 21  # 半磅，即五号 10.5pt


# ---------------------------------------------------------------- 单元格写入工具
def _reset_cell(cell):
    """清空单元格段落，返回原第一段的段落属性作为格式基准。"""
    ps = cell._tc.findall(qn("w:p"))
    base = ps[0].find(qn("w:pPr")) if ps else None
    base = deepcopy(base) if base is not None else OxmlElement("w:pPr")
    for p in ps:
        cell._tc.remove(p)
    return base


def _new_par(cell, base, align=None, indent=False):
    pPr = deepcopy(base)
    for tag in ("w:ind", "w:jc"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if indent:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLineChars"), "200")
        ind.set(qn("w:firstLine"), "420")
        pPr.append(ind)
    if align:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        pPr.append(jc)
    p = OxmlElement("w:p")
    p.append(pPr)
    cell._tc.append(p)
    return Paragraph(p, cell)


def _runs(par, segments, size=BODY_SIZE):
    for text, bold in segments:
        r = par.add_run(text)
        r.bold = bold
        r.font.size = Pt(size / 2)


def fill(cell, blocks, align=None):
    """blocks 中每项：
    ("h", 文本)          加粗小标题
    ("p", 文本)          首行缩进两字的正文
    ("l", 文本)          不缩进的条目
    ("m", [(文本, 粗体)]) 混合格式行
    ("img", 路径, 宽度cm, 图题)
    """
    base = _reset_cell(cell)
    for b in blocks:
        kind = b[0]
        if kind == "h":
            _runs(_new_par(cell, base, align), [(b[1], True)])
        elif kind == "p":
            _runs(_new_par(cell, base, align, indent=True), [(b[1], False)])
        elif kind == "l":
            _runs(_new_par(cell, base, align or "left"), [(b[1], False)])
        elif kind == "m":
            _runs(_new_par(cell, base, align), b[1])
        elif kind == "img":
            _, path, width, caption = b
            _new_par(cell, base, "center").add_run().add_picture(str(path), width=Cm(width))
            _runs(_new_par(cell, base, "center"), [(caption, False)], size=18)


def text(cell, s, align="keep", size=BODY_SIZE):
    """短文本单元格：沿用原对齐方式。"""
    base = _reset_cell(cell)
    jc = base.find(qn("w:jc"))
    keep = jc.get(qn("w:val")) if jc is not None else None
    _runs(_new_par(cell, base, keep if align == "keep" else align), [(s, False)], size=size)


def fix_layout(d):
    """允许表格行跨页断开（否则超过一页的内容会被截掉），并让各节标题行与下一行同页。"""
    for t in d.tables[1:]:
        for row in t.rows:
            trPr = row._tr.trPr
            if trPr is not None:
                for el in trPr.findall(qn("w:cantSplit")):
                    trPr.remove(el)
    T = d.tables
    heads = [(i, 0) for i in range(1, len(T))] + [(3, 1), (8, 1), (10, 1), (10, 6), (11, 1)]
    for ti, ri in heads:
        for cell in T[ti].rows[ri].cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- 内容
SEC1 = [
    ("h", "1. 课题来源"),
    ("p", "本课题来源于油气勘探开发中的实际需求，并结合本学期《软件工程实践》课程实训开展。地下介质速度场是时深转换、"
          "叠前深度偏移成像、构造解释与储层预测的基础参数，其精度直接影响构造落实程度、井位部署与储量评估。"
          "常规速度建模方法各有局限：由叠加速度或层析反演得到的速度场横向连续，但分辨率偏低；全波形反演（FWI）精度高，"
          "但计算代价大、依赖初始模型且易陷入周波跳跃；测井速度（由声波时差换算）垂向分辨率高、精度可靠，"
          "却只分布在稀疏井点，井间只能依靠插值外推。如何把横向连续的地震信息与纵向精确的测井信息有效融合，"
          "是速度建模中的核心问题。"),
    ("p", "近年来，以卷积神经网络为代表的深度学习方法在地震反演与速度建模中展现出强大的非线性映射能力"
          "[2]-[5]，为井震联合建模提供了新思路。本课题拟以测井速度作为井点约束、地震数据作为横向结构引导，"
          "研究基于深度学习的井震联合速度场建模方法，并将其开发为一套可交互使用的桌面软件系统。"),
    ("h", "2. 应用场景与目标用户"),
    ("p", "应用场景：（1）勘探阶段构造解释中的时深转换与深度域成图；（2）为叠前深度偏移、FWI 等处理流程快速构建初始速度模型；"
          "（3）开发阶段储层预测前的背景速度建模；（4）井震联合建模方法的教学演示与科研对比实验。"),
    ("p", "目标用户：油田研究院及物探处理解释人员（主要用户），需要导入工区地震与测井数据，快速获得可靠的速度场并评估其精度；"
          "地球物理方向的科研人员与学生（次要用户），需要在统一平台上对比传统插值方法与深度学习方法的建模效果。"),
    ("h", "3. 现有系统或前期版本情况"),
    ("p", "本课题为全新开发，无前期版本。与本课题相关的现有工具情况如下：（1）商业解释软件（如 Petrel、Jason、GeoEast 等）"
          "提供基于层位约束与克里金插值的速度建模模块，功能成熟，但价格昂贵、算法封闭，难以嵌入深度学习方法，"
          "且高度依赖人工层位解释；（2）开源工具 VelRecover[8] 基于 Python 实现了由稀疏速度拾取点插值生成二维速度剖面的图形界面，"
          "提供线性、对数、RBF、高斯平滑等插值方法，但只利用速度拾取点，既不使用测井资料，也不利用地震振幅中的横向结构信息，"
          "且未引入机器学习方法；（3）学术界已有 InversionNet[3]、OpenFWI[7] 等深度学习速度建模研究代码，"
          "但多为脚本形式，主要面向炮集到速度的端到端映射，缺少测井约束机制，也缺少面向业务人员的数据管理、可视化与精度评估界面。"),
    ("p", "本次开发目标：研究一种测井约束的深度学习井震联合速度建模方法，实现集“数据管理—预处理与井震标定—初始模型构建—"
          "AI 速度建模—可视化评估—成果导出”于一体的开源桌面系统；首先在真值已知的合成模型上验证流程，"
          "以传统插值方法作为基线，并通过盲井检验评估方法的可靠性。"),
]

FUNCS = [
    ("数据管理",
     "新建、打开、保存工程；导入 SEG-Y、NPY、BIN 格式的地震剖面与速度模型，导入 LAS/CSV 格式测井曲线及井位坐标；"
     "浏览数据列表，查看道数、采样点数、采样间隔、网格范围等元信息，并进行格式与维度校验。", "高"),
    ("数据预处理与井震标定",
     "测井曲线异常值剔除与缺失段插补；声波时差 DT 换算为纵波速度 Vp；按地震尺度对测井进行粗化（滑动平均/Backus 平均）；"
     "将井位投影到地震网格并提取井旁道；地震振幅归一化与网格重采样。", "高"),
    ("初始速度模型构建",
     "基于井点速度进行反距离加权（IDW）、普通克里金、径向基函数（RBF）插值，生成低频背景速度模型，"
     "既作为 AI 网络的输入之一，也作为精度对比的传统基线方法。", "中"),
    ("AI 井震联合速度建模",
     "由已知模型或工区数据生成训练样本（随机抽取伪井、切片、数据增强）；配置网络结构与超参数；"
     "在后台线程中训练测井约束 U-Net[6] 并实时显示损失曲线；加载训练好的模型对目标剖面预测速度场。", "高"),
    ("结果可视化与精度评估",
     "速度剖面、地震剖面、误差剖面的显示与色标调节；井位处预测曲线与实测曲线对比；留一法盲井检验；"
          "计算 RMSE、MAE、MAPE、SSIM、R² 等指标，并在固定测试集上与传统插值方法对比。", "高"),
    ("成果导出与报告",
     "将速度场导出为 SEG-Y/NPY/BIN/CSV 格式；导出剖面图与对比图（PNG/PDF）；自动生成精度评估报告；"
     "管理模型权重与训练记录。", "中"),
]

NONFUNC = [
    ("m", [("（1）性能：", True), ("对 1601×401 规模的二维剖面，GPU 环境下单次预测不超过 5 s，CPU 环境下不超过 60 s；"
                                  "训练支持 GPU 加速，无 GPU 时自动回退到 CPU；训练、插值等耗时任务放在后台线程执行，界面保持响应。", False)]),
    ("m", [("（2）易用性：", True), ("按“导入—预处理—建模—评估—导出”的流程组织界面；关键参数提供默认值与说明提示；"
                                   "操作失败时给出明确的中文提示信息。", False)]),
    ("m", [("（3）可靠性：", True), ("导入时校验文件格式、网格维度与数值范围，防止非法数据进入计算；训练过程定期保存检查点，"
                                   "异常中断后可恢复；运行日志写入文件，便于排查问题。", False)]),
    ("m", [("（4）安全性：", True), ("工区原始数据均在本地处理且不上传；代码和脱敏示例数据可通过 GitHub 管理；工程文件与原始数据分开存放，处理过程不修改原始数据。", False)]),
    ("m", [("（5）兼容性：", True), ("支持 Windows 10/11 与 Python 3.10 及以上版本；兼容 SEG-Y rev1、LAS 2.0 等行业标准格式。", False)]),
    ("m", [("（6）可维护性与可扩展性：", True), ("采用分层架构与模块化设计，算法层与界面层解耦；代码遵循 PEP 8 规范，"
                                         "核心模块单元测试覆盖率不低于 70%；网络模型与插值方法通过统一接口注册，"
                                         "便于后续接入 Transformer、扩散模型等新方法。", False)]),
]

SEC3 = [
    ("h", "1. 总体架构"),
    ("p", "系统采用单机桌面应用形态，内部按分层架构组织，自上而下分为表示层、业务逻辑层、算法层和数据访问层。"
          "表示层基于 PySide6 实现图形界面与交互；业务逻辑层负责流程编排与任务调度，训练等耗时任务在后台线程中执行；"
          "算法层封装测井处理、传统插值、深度学习模型与精度指标等核心算法，不依赖界面，可单独进行单元测试与命令行调用；"
          "数据访问层统一负责各类数据文件的读写和 SQLite 工程库的访问。各层之间只通过接口单向依赖。"),
    ("img", FIG / "architecture.png", 14.5, "图1  系统总体架构图"),
    ("h", "2. 功能模块划分"),
    ("p", "系统划分为数据管理、预处理与标定、初始模型构建、AI 速度建模、可视化与评估、成果导出六个功能模块，"
          "与需求分析中的 F1～F6 一一对应，功能结构如图2所示。"),
    ("img", FIG / "function_tree.png", 15.0, "图2  系统功能结构图"),
    ("p", "核心算法的技术路线如图3所示：测井曲线经清洗、DT→Vp 换算并粗化到地震尺度后，由井点插值得到低频背景速度模型；"
          "网络采用 U-Net[6] 编码—解码结构，以归一化地震振幅、低频背景速度和井位掩码三个通道作为输入，以速度场作为输出；"
          "损失函数由全场均方误差与井位处的约束项加权组成（L=MSE+λ·MSE_well），使预测结果在井点处与实测速度保持一致。"
          "监督训练阶段主要在真值已知的 Marmousi/Marmousi2 合成模型上随机抽取若干道作为“伪井”并随机切片生成样本；"
          "真实测井数据用于预处理验证、盲井测试或推理演示；测试阶段在未参与训练的区域和盲井上，"
          "与传统插值基线进行定量对比。"),
    ("img", FIG / "roadmap.png", 15.0, "图3  测井约束的井震联合速度建模技术路线"),
    ("h", "3. 数据设计"),
    ("m", [("数据来源：", True), ("① Marmousi2 模型[1]：纵波速度模型（1601×401 网格，速度 1.03～4.70 km/s）及合成地震剖面（2721×701）。"
                                "两类数据在训练前按坐标范围、采样间隔和网格尺寸进行裁剪与重采样，"
                                "真值已知，用于训练与定量评估；② Marmousi 原始速度模型（2301×751 网格，速度 1.5～5.5 km/s），"
                                "用于跨模型泛化测试；③ 新西兰 Taranaki 盆地整理测井数据集[9]（含声波、密度、伽马等曲线及井位坐标），"
                                "用于测井读取与预处理模块的开发和验证。", False)]),
    ("m", [("文件格式：", True), ("地震与速度数据体采用 SEG-Y、NPY、BIN（float32）；测井采用 LAS 2.0/CSV；"
                                "网络权重采用 PyTorch .pt 格式；工程配置采用 JSON。大体量数据体以文件形式存储，数据库只保存元数据与路径。", False)]),
    ("m", [("数据库表结构（SQLite）：", True), ("", False)]),
    ("l", "project（工程）：project_id 主键，name，root_path，description，created_at"),
    ("l", "dataset（数据体）：dataset_id 主键，project_id 外键，type（地震/速度/结果），file_path，format，n_traces，n_samples，dx，dz"),
    ("l", "well（井）：well_id 主键，project_id 外键，name，x，y，trace_index（投影道号），kb（补心海拔）"),
    ("l", "well_log（测井曲线）：log_id 主键，well_id 外键，curve_name，unit，file_path，depth_start，depth_end，step"),
    ("l", "model_run（建模任务）：run_id 主键，project_id 外键，method，params_json，weight_path，status，result_dataset_id 外键，started_at，finished_at"),
    ("l", "evaluation（评估记录）：eval_id 主键，run_id 外键，metric，value，blind_well_id 外键（可空）"),
    ("p", "实体关系：一个工程包含多个数据体和多口井（1:N）；一口井有多条测井曲线（1:N）；一个工程可发起多次建模任务（1:N）；"
          "每次建模任务生成一个结果数据体（1:1）和多条评估记录（1:N）。"),
    ("h", "4. 界面设计"),
    ("p", "主窗口采用经典的“工程树 + 绘图区 + 参数面板”布局：顶部为菜单栏与工具栏；左侧为工程树，列出数据体、井和建模任务；"
          "中部为多标签页绘图区，显示地震剖面、速度剖面、误差剖面与井曲线对比；右侧为参数面板，随当前步骤切换内容；"
          "底部为日志与进度栏。"),
    ("p", "主要界面包括：工程与数据导入界面、预处理界面（测井曲线编辑、DT→Vp 换算、粗化前后对比）、建模界面（选择方法、"
          "设置超参数、实时损失曲线）、评估界面（剖面对比、井曲线对比、指标表）和导出界面。典型交互流程为：新建工程 → 导入数据 → "
          "预处理与井震标定 → 构建初始模型 → 训练或加载 AI 模型并预测 → 结果评估与对比 → 导出成果。"),
]

SEC5 = [
    ("m", [("开发模型：□ 瀑布模型   ☑ 增量模型   □ 敏捷开发（Scrum）   □ 其他：", False)]),
    ("m", [("选择理由：", True), ("系统功能可以自然分解为相对独立的模块，而最核心、风险最大的是 AI 建模算法能否达到预期精度。"
                                "采用增量模型[10]，可以先交付包含“数据导入—传统插值—AI 建模”核心链路的可运行版本，尽早验证算法可行性，"
                                "再逐步补齐预处理、评估、导出与界面功能，从而降低算法效果不达预期带来的整体风险。"
                                "每个增量内部按“设计—编码—测试”的小迭代推进。", False)]),
    ("m", [("迭代计划：", True), ("", False)]),
    ("l", "第一增量（核心链路，v0.1）：F1 数据导入、F3 传统插值、F4 测井约束 U-Net 训练与预测（先以命令行形式实现），"
          "在 Marmousi2 上跑通并验证方法可行性。"),
    ("l", "第二增量（功能完善，v0.2）：搭建 PySide6 主界面，完成 F2 预处理与井震标定、F5 可视化与精度评估、F6 成果导出与报告。"),
    ("l", "第三增量（交付版本，v1.0）：性能优化与异常处理，完成系统测试与缺陷修复、打包发布与文档编写。"),
    ("m", [("需求变更管理：", True), ("需求按功能编号（F1～F6）登记在 GitHub Issues 中；提出变更时先评估其对进度和已完成模块的影响，"
                                  "经与指导教师确认后纳入下一个增量；每个增量结束时同步更新需求规格说明书，并打版本标签。", False)]),
]

TEAM = [
    ("本人", "F1 数据管理\nF2 预处理与标定", "数据读写模块（segyio/lasio）、SQLite 工程库设计；测井清洗、DT→Vp 换算与粗化、井旁道提取"),
    ("本人", "F3 初始模型构建\nF4 AI 速度建模", "传统插值基线实现；训练样本生成、测井约束 U-Net 设计、训练调参与预测接口"),
    ("本人", "F5 可视化与评估\nF6 成果导出", "剖面与井曲线绘图、盲井检验与指标计算；数据与图件导出、评估报告生成"),
    ("本人", "界面集成、测试与文档", "PySide6 主界面与交互流程；单元/集成/系统测试；需求与设计文档、用户手册与答辩材料"),
]

SEC6 = [
    ("l", "1. 测试类型：☑ 单元测试   ☑ 集成测试   ☑ 系统测试   ☑ 性能测试   ☑ 验收测试"),
    ("l", "2. 测试方法与工具："),
    ("p", "白盒测试：对算法层函数（数据读写、DT→Vp 换算、插值、指标计算等）使用 pytest 编写单元测试，并用 pytest-cov 统计覆盖率，"
          "目标不低于 70%。黑盒测试：按需求规格对各功能模块进行等价类划分与边界值测试（如空文件、维度不匹配、单井、无井等情况）。"
          "界面测试：使用 pytest-qt 模拟用户操作。性能测试：记录不同剖面规模下的训练与预测耗时、内存与显存占用。"
          "算法验收：在合成模型上以真实速度为标准计算误差指标，并进行留一法盲井检验。"),
    ("l", "3. 测试数据来源："),
    ("p", "Marmousi2 速度模型及其合成地震剖面（真值已知，用于定量评估，训练前统一网格）；Marmousi 原始速度模型（用于跨模型泛化测试）；"
          "Taranaki 盆地真实测井数据（文件完整性校验通过后，用于测井读取与预处理模块测试）；人为构造的异常文件（格式错误、缺值、维度不符等，用于健壮性测试）。"),
]

CASES = [
    ("测井数据导入与校验",
     "分别导入标准 LAS 2.0 文件、缺少深度列的 CSV 文件、非测井格式的文本文件",
     "标准文件正确解析出曲线名称、单位与深度范围；异常文件给出明确错误提示，程序不崩溃"),
    ("DT→Vp 换算与测井粗化",
     "输入含 -999.25 空值和尖峰异常值的已知声波时差曲线，执行换算与粗化",
     "按 Vp = 304800/DT（DT 单位 μs/ft）换算正确，相对误差 < 0.1%；空值被标记、异常值被剔除，粗化曲线平滑且无深度偏移"),
    ("AI 井震联合速度建模",
     "在 Marmousi2 剖面上选取 5 口伪井作为输入、另取 2 口作为盲井，执行训练与预测",
     "记录预测速度场的 MAPE、RMSE 等指标；以 MAPE ≤ 8%、盲井处 RMSE 相比 IDW/克里金基线降低 20% 作为预期目标，"
     "最终以固定测试集上的实测结果为准；训练期间界面不卡顿"),
    ("结果导出一致性",
     "将预测速度场分别导出为 SEG-Y 与 NPY 文件，再重新导入",
     "重新导入的数据与原结果维度一致，数值相对误差 < 1e-6"),
]

SEC7 = [
    ("m", [("（1）地球物理专业知识不足。", True), ("井震标定、时深转换、测井粗化等环节涉及较多专业知识。应对：系统阅读速度建模相关文献，"
                                              "先在真值已知的合成模型上开发和验证，关键环节及时向指导教师请教。", False)]),
    ("m", [("（2）井资料稀疏、训练样本不足。", True), ("实际工区井数有限，难以支撑深度网络训练。应对：在 Marmousi/Marmousi2 等标准模型上"
                                               "随机抽取伪井、随机切片，并进行翻转、振幅扰动等数据增强；网络输入中加入井插值低频背景，降低学习难度。", False)]),
    ("m", [("（3）井震尺度不匹配。", True), ("测井采样间隔（约 0.15 m）远小于地震分辨率，直接使用会引入高频噪声。应对：采用滑动平均/"
                                         "Backus 平均将测井粗化到地震尺度，并统一到同一深度网格。", False)]),
    ("m", [("（4）模型泛化能力与结果可信度。", True), ("网络可能在训练模型上过拟合。应对：采用盲井检验与跨模型测试评估泛化能力；"
                                               "在损失函数中加入井点约束项，使井位处结果与实测基本一致；保留传统插值结果作为对照。此外，公开渠道难以获得同一工区配套的实测地震与测井数据，真实数据上的效果只作演示性验证，定量结论以合成模型为准。", False)]),
    ("m", [("（5）计算资源有限。", True), ("个人电脑显存有限。应对：采用切片训练与混合精度，控制网络规模；支持 CPU 回退，"
                                        "并提供预训练权重供直接加载。", False)]),
    ("m", [("（6）数据获取与进度风险。", True), ("部分公开数据集体量较大（如 Taranaki 测井数据约 229 MB），下载易中断；算法调参耗时难以预估。"
                                             "应对：提前完整下载并校验数据；按增量开发，优先保证核心链路可运行，每周对照进度表检查，"
                                             "必要时压缩非核心功能（如报告自动生成）。", False)]),
]

SCHEDULE = [
    ("2026.09.28～2026.10.11", "需求调研与分析，完成开题任务书"),
    ("2026.10.12～2026.10.25", "概要设计与详细设计（架构、模块、数据库、界面）"),
    ("2026.10.26～2026.11.22", "编码实现：第一轮迭代，完成核心功能（数据导入、传统插值、测井约束 U-Net 训练与预测）"),
    ("2026.11.23～2026.12.13", "编码实现：第二轮迭代，完善其余功能（预处理与标定、可视化评估、成果导出、图形界面）"),
    ("2026.12.14～2026.12.27", "第三轮迭代：性能优化与异常处理，系统测试、缺陷修复与打包发布"),
    ("2026.12.28～2027.01.10", "编写文档，整理成果，准备答辩"),
]

DELIVERABLES = [
    "☑ 软件需求规格说明书",
    "☑ 软件设计说明书（概要设计、详细设计）",
    "☑ 完整源代码及 README（环境配置、编译运行步骤）",
    "☑ 可执行程序或部署包",
    "☑ 测试报告",
    "☑ 用户使用手册",
    "☑ 课程设计报告",
    "☑ 答辩 PPT 及演示视频",
]

REFS = [
    "[1] Martin G S, Wiley R, Marfurt K J. Marmousi2: An elastic upgrade for Marmousi[J]. The Leading Edge, 2006, 25(2): 156-166.",
    "[2] Yang F, Ma J. Deep-learning inversion: A next-generation seismic velocity model building method[J]. Geophysics, 2019, 84(4): R583-R599.",
    "[3] Wu Y, Lin Y. InversionNet: An efficient and accurate data-driven full waveform inversion[J]. IEEE Transactions on Computational Imaging, 2020, 6: 419-433.",
    "[4] Araya-Polo M, Jennings J, Adler A, et al. Deep-learning tomography[J]. The Leading Edge, 2018, 37(1): 58-66.",
    "[5] Das V, Pollack A, Wollner U, et al. Convolutional neural network for seismic impedance inversion[J]. Geophysics, 2019, 84(6): R869-R880.",
    "[6] Ronneberger O, Fischer P, Brox T. U-Net: Convolutional networks for biomedical image segmentation[C]//Medical Image Computing and Computer-Assisted Intervention (MICCAI). Cham: Springer, 2015: 234-241.",
    "[7] Deng C, Feng S, Wang H, et al. OpenFWI: Large-scale multi-structural benchmark datasets for full waveform inversion[C]//Advances in Neural Information Processing Systems 35 (NeurIPS 2022), Datasets and Benchmarks Track, 2022.",
    "[8] Pertuz A, Benito M I, Llanes Estrada P, et al. VelRecover: An interactive Python tool for building velocity models from sparse velocity picks in legacy seismic sections[CP/OL]. Zenodo, 2025. DOI: 10.5281/zenodo.15053268.",
    "[9] de Carvalho B W S R, Oliveira M, Avalone M, et al. Taranaki Basin curated well logs[DS/OL]. Zenodo, 2020. DOI: 10.5281/zenodo.3832955.",
    "[10] Sommerville I. Software Engineering[M]. 10th ed. Boston: Pearson, 2016.",
]


def main():
    d = docx.Document(str(TEMPLATE))
    T = d.tables

    # 封面
    text(T[0].rows[0].cells[1], TITLE, size=24)
    date_runs = d.paragraphs[13].runs
    date_runs[1].text, date_runs[3].text, date_runs[5].text = " 2026 ", " 9 ", " 28 "

    # 基本信息
    text(T[1].rows[2].cells[1], TITLE)
    text(T[1].rows[3].cells[1], "单人选题")

    # 一、项目背景
    fill(T[2].rows[1].cells[0], SEC1)

    # 二、功能需求
    for i, (name, desc, prio) in enumerate(FUNCS):
        row = T[3].rows[2 + i]
        text(row.cells[1], name)
        fill(row.cells[2], [("l", desc)])
        text(row.cells[3], {"高": "☑高 □中 □低", "中": "□高 ☑中 □低", "低": "□高 □中 ☑低"}[prio])

    # 非功能需求
    fill(T[4].rows[1].cells[0], NONFUNC)

    # 三、系统设计
    fill(T[5].rows[1].cells[0], SEC3)

    # 四、开发环境
    t6 = T[6]
    text(t6.rows[1].cells[1], "□ C#   □ C++   □ Java   ☑ Python   □ 其他：")
    fill(t6.rows[2].cells[1], [
        ("l", "界面：PySide6（Qt for Python），绘图：Matplotlib / pyqtgraph"),
        ("l", "深度学习：PyTorch；科学计算：NumPy、SciPy、scikit-learn、PyKrige"),
        ("l", "数据读写：segyio（SEG-Y）、lasio（LAS）、pandas；打包：PyInstaller"),
    ])
    fill(t6.rows[3].cells[1], [("l", "SQLite 3（Python 内置 sqlite3 模块），存储工程、数据体、井、建模任务与评估记录等元数据；"
                                     "大体量数据体以文件形式存储")])
    fill(t6.rows[4].cells[1], [("l", "Visual Studio Code（主力开发）、Jupyter Notebook（算法原型实验）")])
    text(t6.rows[5].cells[1], "☑ Windows   □ Linux   □ Web   □ 其他：")
    fill(t6.rows[6].cells[1], [("l", "Git + GitHub，仓库地址：https://github.com/Lx123759/well-seismic-velocity")])

    # 五、开发过程
    fill(T[7].rows[1].cells[0], SEC5)

    # 小组分工
    for i, (who, mod, work) in enumerate(TEAM):
        row = T[8].rows[2 + i]
        text(row.cells[0], who, align="center")
        fill(row.cells[1], [("l", s) for s in mod.split("\n")], align="center")
        fill(row.cells[2], [("l", work)])

    # 六、测试计划
    fill(T[9].rows[1].cells[0], SEC6)
    for i, (func, op, exp) in enumerate(CASES):
        row = T[10].rows[2 + i]
        fill(row.cells[1], [("l", func)])
        fill(row.cells[2], [("l", op)])
        fill(row.cells[3], [("l", exp)])

    # 七、难点
    fill(T[10].rows[7].cells[0], SEC7)

    # 八、进度
    for i, (period, work) in enumerate(SCHEDULE):
        row = T[11].rows[2 + i]
        text(row.cells[1], period)
        fill(row.cells[2], [("l", work)])

    # 九、成果
    fill(T[12].rows[1].cells[0], [("l", s) for s in DELIVERABLES])

    # 十、参考文献
    fill(T[13].rows[1].cells[0], [("l", s) for s in REFS])

    fix_layout(d)
    d.save(str(OUTPUT))
    print("saved:", OUTPUT.name)


if __name__ == "__main__":
    main()
